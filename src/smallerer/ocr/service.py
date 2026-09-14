"""Page-level OCR orchestration: rasterization, caching, resume (spec §5.5).

Key behaviors:
- Per-page decision: OCR only pages with text_layer < 30 chars (ocr_page_threshold)
- Page-level cache: content_hash + page_no + ocr config digest
- Graceful degradation: no backend → warning + skip (unless --ocr only)
- Interrupt resume: cache persists across runs, only missing pages re-OCR
- Max pixels: 4000×4000 per page to avoid Vision API limits
"""

from __future__ import annotations

import hashlib
from pathlib import Path

from ..config import CACHE_DIR_NAME, Config, OcrMode
from ..model import Document, ExtractError, Kind, Line
from . import get_backend, unavailable_reason
from .base import OcrBackend
from .cache import OcrCache

MAX_PIXELS = 4000


def apply(
    doc: Document,
    path: Path,
    kind: Kind,
    cfg: Config,
    content_hash: str,
) -> None:
    """Apply OCR to pages needing it, with caching and backend availability checks.

    Backend unavailability handling (spec §6.3):
    - --ocr only + no backend → ExtractError (exit 2)
    - --ocr auto + no backend → warning + graceful skip
    - macOS: Vision available
    - Linux/Windows: no backend yet, degrades to never

    Cache key: content_hash + page_no + digest(ocr_dpi, max_pixels, backend)
    """
    pages = doc.pages_needing_ocr
    if not pages or cfg.ocr is OcrMode.NEVER:
        return

    backend = get_backend()
    if backend is None:
        if cfg.ocr is OcrMode.ONLY:
            raise ExtractError("no_ocr_backend", unavailable_reason())
        doc.warnings.append(unavailable_reason())
        return

    selected = pages[: cfg.max_ocr_pages]
    if len(pages) > len(selected):
        doc.partial = True
        doc.warnings.append(
            f"需 OCR {len(pages)} 页，超过上限 {cfg.max_ocr_pages}，"
            f"只处理前 {len(selected)} 页"
        )

    assert cfg.out is not None
    cache = OcrCache(cfg.out / CACHE_DIR_NAME / "ocr")
    digest = _cache_digest(cfg, backend)

    for page_number in selected:
        blocks = cache.get(content_hash, page_number, digest)
        if blocks is None:
            image_payloads = _render(path, kind, page_number, cfg)
            blocks = []
            for payload in image_payloads:
                blocks.extend(backend.recognize(payload))
            blocks.sort(key=lambda block: block.y0 if block.y0 is not None else 1.0)
            cache.put(content_hash, page_number, digest, backend.name, blocks)

        page = next((p for p in doc.pages if p.number == page_number), None)
        if page is None:
            continue
        page.lines = [
            Line(text=block.text, y0=block.y0, y1=block.y1)
            for block in blocks
            if block.text.strip()
        ]
        page.height = 1.0
        page.source = "ocr"
        doc.ocr_pages.append(page_number)
        doc.ocr_chars += sum(len(block.text) for block in blocks)
        if not blocks:
            doc.warnings.append(f"第 {page_number} 页 OCR 没有识别出任何文字")

    doc.ocr_backend = backend.name
    if doc.ocr_pages:
        doc.extractor = f"{doc.extractor}+{backend.name}"


def _cache_digest(cfg: Config, backend: OcrBackend) -> str:
    raw = (
        f"{cfg.content_digest()}\0{backend.name}\0"
        f"dpi={cfg.ocr_dpi}\0max_pixels={MAX_PIXELS}"
    ).encode()
    return hashlib.blake2b(raw, digest_size=12).hexdigest()


def _render(path: Path, kind: Kind, page_number: int, cfg: Config) -> list[bytes]:
    if kind is Kind.PDF:
        return [_render_pdf(path, page_number, cfg.ocr_dpi)]
    if kind is Kind.IMAGE:
        return [path.read_bytes()]
    if kind is Kind.PPTX:
        payloads = _pptx_images(path, page_number)
        if payloads:
            return payloads
        raise ExtractError(
            "pptx_render_unavailable",
            f"第 {page_number} 张幻灯片没有可直接识别的内嵌图片；"
            "第一版不依赖 LibreOffice，无法渲染整张幻灯片",
        )
    raise ExtractError("ocr_unsupported", f"{kind.value} 不支持 OCR")


def _render_pdf(path: Path, page_number: int, dpi: int) -> bytes:
    import pymupdf

    try:
        with pymupdf.open(path) as doc:
            page = doc.load_page(page_number - 1)
            scale = dpi / 72.0
            max_dimension = max(page.rect.width, page.rect.height) * scale
            if max_dimension > MAX_PIXELS:
                scale *= MAX_PIXELS / max_dimension
            pixmap = page.get_pixmap(matrix=pymupdf.Matrix(scale, scale), alpha=False)
            return pixmap.tobytes("png")
    except Exception as exc:
        raise ExtractError("ocr_rasterize_failed", f"第 {page_number} 页光栅化失败：{exc}") from exc


def _pptx_images(path: Path, page_number: int) -> list[bytes]:
    """纯图 PPTX 的第一版路径：识别该页所有内嵌图片，按形状位置排序。

    这不会渲染矢量图、SmartArt 或主题装饰；因此只在文字极少的页触发，而且不足会明确报错。
    """
    try:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE

        deck = Presentation(str(path))
        slide = deck.slides[page_number - 1]
        pictures = [
            shape
            for shape in slide.shapes
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
        ]
        pictures.sort(key=lambda shape: (shape.top, shape.left))
        return [bytes(shape.image.blob) for shape in pictures]
    except Exception as exc:
        raise ExtractError(
            "pptx_image_read_failed",
            f"第 {page_number} 张幻灯片读取内嵌图片失败：{exc}",
        ) from exc

