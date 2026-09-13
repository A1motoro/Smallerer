"""PPTX extractor（spec §6.2）：形状文字 + 演讲者备注；纯图页标记待 OCR。"""

from __future__ import annotations

import re
from pathlib import Path

from ..config import Config
from ..model import Document, ExtractError, Kind, Line, Page, TextLayer

_WS_RE = re.compile(r"\s+")
_SPARSE_CHARS = 200


def _nonspace(text: str) -> int:
    return len(_WS_RE.sub("", text))


def _table_rows(shape) -> list[str]:
    table = shape.table
    rows = [[" ".join(cell.text.split()) for cell in row.cells] for row in table.rows]
    if not rows or max(len(r) for r in rows) < 2:
        return []
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    out = ["| " + " | ".join(rows[0]) + " |", "|" + "---|" * width]
    for row in rows[1:]:
        out.append("| " + " | ".join(row) + " |")
    return out


def extract(path: Path, cfg: Config) -> Document:
    try:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE
    except ImportError as exc:
        raise ExtractError("missing_dependency", f"python-pptx 不可用：{exc}") from exc

    try:
        deck = Presentation(str(path))
    except Exception as exc:
        raise ExtractError("corrupt", f"PPTX 无法打开：{exc}") from exc

    out = Document(kind=Kind.PPTX, extractor="python-pptx", page_label="Slide")
    out.total_pages = len(deck.slides)
    page_chars: list[int] = []

    for index, slide in enumerate(deck.slides, start=1):
        lines: list[Line] = []
        pictures = 0

        for shape in _ordered(slide.shapes):
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                pictures += 1
                continue
            if getattr(shape, "has_table", False):
                rows = _table_rows(shape)
                if rows:
                    out.has_tables = True
                    lines.extend(Line(text=row) for row in rows)
                    lines.append(Line(text=""))
                continue
            if not getattr(shape, "has_text_frame", False):
                continue
            for paragraph in shape.text_frame.paragraphs:
                text = " ".join("".join(run.text for run in paragraph.runs).split())
                if not text:
                    continue
                bullet = "  " * max(paragraph.level, 0)
                lines.append(Line(text=f"{bullet}{text}" if paragraph.level else text))
            lines.append(Line(text=""))

        notes = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        if notes:
            lines.append(Line(text=""))
            lines.append(Line(text="**演讲者备注**"))
            lines.append(Line(text=""))
            for para in notes.split("\n"):
                cleaned = " ".join(para.split())
                if cleaned:
                    lines.append(Line(text=cleaned))

        body = "".join(line.text for line in lines)
        chars = _nonspace(body)
        page_chars.append(chars)
        out.pages.append(Page(number=index, lines=lines))
        if chars < cfg.ocr_page_threshold:
            out.pages_needing_ocr.append(index)
            if pictures:
                out.image_heavy_pages.append(index)

    out.text_layer = _classify(page_chars, cfg.ocr_page_threshold)
    return out


def _ordered(shapes):
    """按位置排出阅读顺序；没有位置信息的形状排在最后。"""

    def key(shape):
        top = getattr(shape, "top", None)
        left = getattr(shape, "left", None)
        return (top if top is not None else 1 << 40, left if left is not None else 0)

    return sorted(shapes, key=key)


def _classify(page_chars: list[int], threshold: int) -> TextLayer:
    if not page_chars:
        return TextLayer.NONE
    empty = sum(1 for c in page_chars if c < threshold)
    rich = sum(1 for c in page_chars if c >= _SPARSE_CHARS)
    if empty == len(page_chars):
        return TextLayer.NONE
    if empty:
        return TextLayer.MIXED
    if rich == len(page_chars):
        return TextLayer.RICH
    return TextLayer.SPARSE
