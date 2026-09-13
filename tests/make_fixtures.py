"""生成测试 fixture。运行结果不提交 git，避免任何版权材料。"""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

from PIL import Image
from reportlab.pdfgen import canvas

HERE = Path(__file__).parent
OUT = HERE / "fixtures" / "generated"


def text_pdf(path: Path) -> None:
    doc = canvas.Canvas(str(path))
    for page in range(1, 4):
        doc.setFont("Helvetica", 10)
        doc.drawString(72, 810, "PHYS 201 — Fall 2026")
        doc.setFont("Helvetica-Bold", 18)
        doc.drawString(72, 760, f"Chapter {page}: Context Extraction")
        doc.setFont("Helvetica", 11)
        y = 720
        for index in range(12):
            doc.drawString(
                72,
                y,
                f"This is paragraph {index + 1} on page {page}. "
                "It contains enough searchable text for quality checks.",
            )
            y -= 24
        doc.setFont("Helvetica", 9)
        doc.drawString(280, 25, f"Page {page}")
        doc.showPage()
    doc.save()


def image_pdf(path: Path) -> None:
    image = Image.new("RGB", (1200, 1600), "white")
    buffer = BytesIO()
    image.save(buffer, "PNG")
    payload = buffer.getvalue()

    doc = canvas.Canvas(str(path), pagesize=(600, 800))
    for _ in range(2):
        doc.drawInlineImage(Image.open(BytesIO(payload)), 0, 0, 600, 800)
        doc.showPage()
    doc.save()


def encrypted_pdf(source: Path, target: Path) -> None:
    import pymupdf

    src = pymupdf.open(source)
    src.save(
        target,
        encryption=pymupdf.PDF_ENCRYPT_AES_256,
        owner_pw="owner",
        user_pw="secret",
    )
    src.close()


def docx_fixture(path: Path) -> None:
    from docx import Document

    doc = Document()
    doc.add_heading("Generated Document", level=1)
    doc.add_paragraph(
        "This paragraph contains enough text to verify that DOCX extraction "
        "preserves headings and paragraph anchors."
    )
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "alpha"
    table.cell(1, 1).text = "42"
    doc.save(path)


def pptx_fixture(path: Path) -> None:
    from pptx import Presentation

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "Generated Slide"
    slide.placeholders[1].text = (
        "This is generated slide content.\n"
        "It is long enough to exercise the PPTX extractor and quality check."
    )
    slide.notes_slide.notes_text_frame.text = "Speaker note for the generated slide."
    deck.save(path)


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    text_pdf(OUT / "text.pdf")
    image_pdf(OUT / "image-only.pdf")
    encrypted_pdf(OUT / "text.pdf", OUT / "encrypted.pdf")
    docx_fixture(OUT / "sample.docx")
    pptx_fixture(OUT / "sample.pptx")
    (OUT / "notes.txt").write_text("Already plain text.\n", encoding="utf-8")
    (OUT / "unknown.bin").write_bytes(b"\x00\x01\x02\x03")
    (OUT / "corrupt.pdf").write_bytes(b"%PDF-1.7\nnot actually a PDF")


if __name__ == "__main__":
    main()

